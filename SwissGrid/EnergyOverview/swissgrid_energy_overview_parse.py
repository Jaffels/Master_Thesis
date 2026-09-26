#!/usr/bin/env python3
"""
Swissgrid Energy Overview parser (EnergieUebersichtCH-YYYY.xls / .xlsx).

Run from the thesis root:
    python Swissgrid/EnergyOverview/swissgrid_energy_overview_parse.py
    python Swissgrid/EnergyOverview/swissgrid_energy_overview_parse.py --years 2026 --force

Needs: pandas, pyarrow, openpyxl (.xlsx), xlrd (.xls)

SAFETY: input files are opened read-only. The script writes ONLY under the
output directory (default Swissgrid/EnergyOverview/Data/) and never deletes
anything. Existing year files are skipped unless --force is given; with
--force they are replaced atomically (write to temp file, then rename).

Source layout (verified on 2019 .xls, 2020 .xlsx, 2025 .xlsx):
  Sheet "Zeitreihen0h15", LONG format:
    row 0 : variable names "German\\nEnglish" (col 0 empty)
    row 1 : "Zeitstempel" + units (kWh, Euro/MWh)
    row 2+: timestamp | 64 values
  Sheet "Zeitreihen1h00" (until 2024): vertical grid load, MW, hourly.
  Timestamps are local wall-clock (Europe/Zurich):
    - up to 2024 (at least 2019/2020): label = interval END, Excel datetime
      (first value 00:15). On the autumn DST day the last CEST interval is
      labelled 03:00 instead of 02:00 -> labels cannot be localised directly.
    - 2025+: label = interval START, text "dd.mm.yyyy HH:MM" (first 00:00).
  Rows are strictly sequential 15-min steps (92/100 rows on DST days), so
  the parser builds the UTC index from the first timestamp + row position
  and then CHECKS every source label against it. Mismatches are allowed only
  on DST switch days; anything else (gap, duplicate) fails that file.

Output (all timestamps = interval START, datetime64[us, Europe/Zurich]):
  Data/qh/<year>.parquet              15-min, fixed 64-column schema (QH_SCHEMA);
                                      NaN where a variable was not yet published.
                                      Load all years: pd.read_parquet("Swissgrid/EnergyOverview/Data/qh/")
  Data/vertical_load_1h/<year>.parquet hourly vertical load (if sheet exists)
  Data/variables.csv                  column -> German/English label, unit
  Data/_parse_manifest.csv            one row per (file, sheet)
Units are kept as published: energy in kWh per interval, prices in EUR/MWh.
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import sys
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd

TZ = "Europe/Zurich"
FILE_RE = re.compile(r"EnergieUebersichtCH[-_](\d{4})\.(xlsx?|XLSX?)$")
QH_SHEET = "Zeitreihen0h15"
H_SHEET = "Zeitreihen1h00"
TEXT_TS_FORMATS = ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M")

# ---------------------------------------------------------------------------
# Column naming: keyed on the normalised ENGLISH part of the header.
# secondary control = aFRR, tertiary control = mFRR.
# ---------------------------------------------------------------------------
SYSTEM_MAP = {
    "total energy consumed by end users in the swiss controlblock": "end_user_consumption",
    "total energy production swiss controlblock": "production",
    "total energy consumption swiss controlblock": "consumption",
    "net outflow of the swiss transmission grid": "net_outflow_tn",
    "grid feed-in swiss transmission grid": "vertical_feedin_tn",
    "positive secundary control energy": "afrr_pos_energy",
    "positive secondary control energy": "afrr_pos_energy",
    "negative secundary control energy": "afrr_neg_energy",
    "negative secondary control energy": "afrr_neg_energy",
    "positive tertiary control energy": "mfrr_pos_energy",
    "negative tertiary control energy": "mfrr_neg_energy",
    "transit": "transit",
    "import": "import",
    "export": "export",
    "average positive secondary control energy prices": "afrr_pos_price",
    "average negative secondary control energy prices": "afrr_neg_price",
    "average positive tertiary control energy prices": "mfrr_pos_price",
    "average negative tertiary control energy prices": "mfrr_neg_price",
    "production across cantons": "prod_cross_canton",
    "consumption across cantons": "cons_cross_canton",
    "production control area ch - foreign territories": "prod_foreign_territories",
    "consumption control area ch - foreign territories": "cons_foreign_territories",
    "vertical load swiss transmission grid": "vertical_load",
}
# Fixed output schema: every year file gets exactly these columns, in this
# order, NaN where the variable was not yet published (2009-13: 20 source
# columns, 2014: 24, 2015+: 64). Names never depend on the source unit text.
QH_SCHEMA = (
    [("end_user_consumption", "kwh"), ("production", "kwh"), ("consumption", "kwh"),
     ("net_outflow_tn", "kwh"), ("vertical_feedin_tn", "kwh"),
     ("afrr_pos_energy", "kwh"), ("afrr_neg_energy", "kwh"),
     ("mfrr_pos_energy", "kwh"), ("mfrr_neg_energy", "kwh")]
    + [(f"xb_{a}_{b}", "kwh") for n in ("at", "de", "fr", "it") for a, b in (("ch", n), (n, "ch"))]
    + [("transit", "kwh"), ("import", "kwh"), ("export", "kwh"),
       ("afrr_pos_price", "eur_mwh"), ("afrr_neg_price", "eur_mwh"),
       ("mfrr_pos_price", "eur_mwh"), ("mfrr_neg_price", "eur_mwh")]
    + [(f"{k}_{c}", "kwh")
       for c in ("ag", "fr", "gl", "gr", "lu", "ne", "so", "sg", "ti", "tg", "vs",
                 "ai_ar", "bl_bs", "be_ju", "sz_zg", "ow_nw_ur", "ge_vd", "sh_zh")
       for k in ("prod", "cons")]
    + [("prod_cross_canton", "kwh"), ("cons_cross_canton", "kwh"),
       ("prod_foreign_territories", "kwh"), ("cons_foreign_territories", "kwh")]
)
H_SCHEMA = [("vertical_load", "mw")]
assert len(QH_SCHEMA) == 64

XB_RE = re.compile(r"^cross border exchange ([a-z]{2})->([a-z]{2})$")
CANTON_RE = re.compile(r"^(production|consumption) cantons? ([a-z ,]+)$")
UNIT_SUFFIX = {"kwh": "kwh", "euro/mwh": "eur_mwh", "eur/mwh": "eur_mwh", "mw": "mw", "mwh": "mwh"}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s).strip().lower())


def column_name(header: str, unit: str) -> tuple[str, bool, str]:
    """Return (base name without unit, mapped_ok, normalised unit suffix)."""
    parts = [p for p in str(header).split("\n") if p.strip()]
    en = _norm(parts[-1]) if parts else ""
    unit_sfx = UNIT_SUFFIX.get(_norm(unit), re.sub(r"[^a-z0-9]+", "_", _norm(unit)).strip("_"))
    base, ok = SYSTEM_MAP.get(en), True
    if base is None:
        m = XB_RE.match(en)
        if m:
            base = f"xb_{m.group(1)}_{m.group(2)}"
        else:
            m = CANTON_RE.match(en)
            if m:
                kind = "prod" if m.group(1) == "production" else "cons"
                cantons = "_".join(c.strip() for c in m.group(2).split(","))
                base = f"{kind}_{cantons}"
            else:
                base = re.sub(r"[^a-z0-9]+", "_", en).strip("_") or "unnamed"
                ok = False
    return base, ok, unit_sfx


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------
def read_sheet(path: Path, sheet: str) -> list[list] | None:
    """Return all rows of a sheet as python lists (None if sheet absent)."""
    if path.suffix.lower() == ".xls":
        import xlrd

        wb = xlrd.open_workbook(str(path), on_demand=True)
        try:
            if sheet not in wb.sheet_names():
                return None
            sh = wb.sheet_by_name(sheet)
            rows = []
            for i in range(sh.nrows):
                row = sh.row_values(i)
                types = sh.row_types(i)
                if types and types[0] == xlrd.XL_CELL_DATE:
                    row[0] = xlrd.xldate.xldate_as_datetime(row[0], wb.datemode)
                rows.append(row)
            return rows
        finally:
            wb.release_resources()
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return None
        return [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()


def find_header(rows: list[list]) -> int:
    """Index of the units row (col 0 == 'Zeitstempel'); names are the row above."""
    for i, r in enumerate(rows[:15]):
        if r and isinstance(r[0], str) and r[0].strip().lower().startswith("zeitstempel"):
            if i == 0:
                raise ValueError("'Zeitstempel' row has no names row above it")
            return i
    raise ValueError("no 'Zeitstempel' header row in first 15 rows")


def to_naive_ts(v) -> pd.Timestamp:
    if isinstance(v, (dt.datetime, pd.Timestamp)):
        return pd.Timestamp(v).round("min")
    if isinstance(v, (int, float)) and not isinstance(v, bool):  # Excel serial
        return (pd.Timestamp("1899-12-30") + pd.to_timedelta(float(v), unit="D")).round("min")
    if isinstance(v, str):
        s = v.strip()
        for fmt in TEXT_TS_FORMATS:
            try:
                return pd.Timestamp(dt.datetime.strptime(s, fmt))
            except ValueError:
                pass
    raise ValueError(f"unparseable timestamp {v!r}")


def to_float(v) -> float:
    if v is None or v == "":
        return np.nan
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    try:
        return float(str(v).replace("'", "").replace(",", "."))
    except ValueError:
        return np.nan


# ---------------------------------------------------------------------------
# Timestamp reconstruction + validation
# ---------------------------------------------------------------------------
def build_index(labels: list[pd.Timestamp], freq_min: int) -> tuple[pd.DatetimeIndex, str, int]:
    """
    Build tz-aware interval-START index from row position and validate labels.
    Returns (index, convention, n_dst_label_mismatches). Raises on real gaps.
    """
    step = pd.Timedelta(minutes=freq_min)
    first = labels[0]
    # Convention: a year file starts at local midnight. Label 00:00 -> start,
    # label 00:<step> -> end. Fall back on time-of-day otherwise.
    tod = first - first.normalize()
    if tod == pd.Timedelta(0):
        conv, first_start = "start", first
    elif tod == step:
        conv, first_start = "end", first - step
    else:
        raise ValueError(f"cannot infer label convention from first timestamp {first}")
    t0 = first_start.tz_localize(TZ)  # never ambiguous: files start in January
    n = len(labels)
    starts_utc = pd.date_range(t0.tz_convert("UTC"), periods=n, freq=step)
    idx = starts_utc.tz_convert(TZ)
    expected = (idx if conv == "start" else idx + step).tz_localize(None)
    got = pd.DatetimeIndex(labels)
    bad = np.flatnonzero(expected != got)
    if len(bad):
        # Allowed only on DST switch days (local date of the interval)
        local_day = idx.tz_localize(None).normalize()
        per_day = pd.Series(local_day).value_counts()
        switch_days = set(per_day.index[per_day != 24 * 60 // freq_min])
        switch_days -= {local_day[0], local_day[-1]}  # partial first/last day is not a DST day
        off = [d for d in set(local_day[bad]) if d not in switch_days]
        if off:
            i = int(bad[0])
            raise ValueError(
                f"{len(bad)} timestamp labels deviate from a continuous {freq_min}-min sequence "
                f"outside DST days (first at row {i}: expected {expected[i]}, found {got[i]}) "
                f"-> gap or duplicate in source"
            )
    return idx.as_unit("us"), conv, len(bad)


# ---------------------------------------------------------------------------
# Per-sheet parse
# ---------------------------------------------------------------------------
def parse_sheet(rows: list[list], freq_min: int, schema: list[tuple[str, str]]):
    h = find_header(rows)
    names, units = rows[h - 1], rows[h]
    ncol = max(len(names), len(units))
    names = list(names) + [None] * (ncol - len(names))
    units = list(units) + [None] * (ncol - len(units))
    keep = [j for j in range(1, ncol) if names[j] not in (None, "")]

    body = [r for r in rows[h + 1 :] if r and r[0] not in (None, "")]
    if not body:
        raise ValueError("no data rows")
    labels = [to_naive_ts(r[0]) for r in body]
    idx, conv, n_mis = build_index(labels, freq_min)

    expected_unit = dict(schema)
    canonical = [f"{b}_{u}" for b, u in schema]
    cols, meta, unmapped, unit_mismatch = {}, [], [], []
    for j in keep:
        base, ok, sfx = column_name(names[j], units[j] or "")
        if ok and base in expected_unit:
            col = f"{base}_{expected_unit[base]}"
            if sfx != expected_unit[base]:
                unit_mismatch.append(f"{col}: source unit {units[j]!r}")
        else:  # unknown variable: keep it (appended after the fixed schema) and flag it
            col, ok = (f"{base}_{sfx}" if sfx else base), False
        root, k = col, 2
        while col in cols:  # duplicate header: keep both, flag the copy
            col, k, ok = f"{root}_{k}", k + 1, False
        if not ok:
            unmapped.append(col)
        cols[col] = np.fromiter((to_float(r[j]) if j < len(r) else np.nan for r in body), float, len(body))
        parts = [p.strip() for p in str(names[j]).split("\n") if p.strip()]
        meta.append({"column": col, "label_de": parts[0] if parts else "",
                     "label_en": parts[-1] if parts else "", "unit": units[j] or ""})
    source_cols = list(cols)
    missing = [c for c in canonical if c not in cols]
    df = pd.DataFrame(cols).reindex(columns=canonical + [c for c in cols if c not in canonical])
    df = df.astype("float64")
    df.insert(0, "timestamp", idx)
    return df, meta, conv, n_mis, unmapped, missing, unit_mismatch, source_cols


def atomic_parquet(df: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".parquet.tmp", dir=path.parent)
    os.close(fd)
    try:
        df.to_parquet(tmp, index=False)
        os.chmod(tmp, 0o644)  # mkstemp creates 0600
        os.replace(tmp, path)
    except BaseException:
        if os.path.exists(tmp):
            os.remove(tmp)  # only our own temp file
        raise


# ---------------------------------------------------------------------------
# Driver
# ---------------------------------------------------------------------------
def discover(input_dir: Path, years: set[int] | None) -> dict[int, Path]:
    found: dict[int, list[Path]] = {}
    for p in sorted(input_dir.iterdir()):
        m = FILE_RE.search(p.name)
        if p.is_file() and m and not p.name.startswith("~$"):
            y = int(m.group(1))
            if years is None or y in years:
                found.setdefault(y, []).append(p)
    out = {}
    for y, ps in found.items():
        if len(ps) > 1:  # e.g. both .xls and .xlsx: take the most recently modified
            ps.sort(key=lambda p: p.stat().st_mtime)
            print(f"  ! {y}: {len(ps)} files, using {ps[-1].name}")
        out[y] = ps[-1]
    return dict(sorted(out.items()))


def resolve_input(root: Path, arg: str | None) -> Path:
    if arg:
        p = Path(arg)
        return p if p.is_absolute() else root / p
    default = root / "Swissgrid/Manual_Download/Balancing"
    if default.is_dir():
        return default
    cands = [d for d in (root / "Swissgrid/Manual_Download").rglob("Balancing") if d.is_dir()]
    cands = [d for d in cands if any(FILE_RE.search(f.name) for f in d.iterdir())]
    if len(cands) == 1:
        return cands[0]
    raise SystemExit(
        "Could not locate the Energy Overview 'Balancing' folder. Pass --input-dir. "
        f"Candidates: {[str(c) for c in cands] or 'none'}"
    )


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--root", default=".", help="thesis root (default: cwd)")
    ap.add_argument("--input-dir", help="folder with EnergieUebersichtCH-YYYY files "
                    "(default: Swissgrid/Manual_Download/Balancing)")
    ap.add_argument("--output-dir", default="Swissgrid/EnergyOverview/Data")
    ap.add_argument("--years", help="comma list / ranges, e.g. 2019,2021-2023")
    ap.add_argument("--force", action="store_true", help="re-parse years that already have output")
    a = ap.parse_args(argv)

    root = Path(a.root).resolve()
    in_dir = resolve_input(root, a.input_dir)
    out_dir = Path(a.output_dir)
    out_dir = out_dir if out_dir.is_absolute() else root / out_dir
    years = None
    if a.years:
        years = set()
        for part in a.years.split(","):
            lo, _, hi = part.partition("-")
            years.update(range(int(lo), int(hi or lo) + 1))

    files = discover(in_dir, years)
    print(f"Input : {in_dir}\nOutput: {out_dir}\nFiles : {len(files)}")
    man_path = out_dir / "_parse_manifest.csv"
    manifest = pd.read_csv(man_path) if man_path.exists() else pd.DataFrame()
    var_path = out_dir / "variables.csv"
    canonical_all = [f"{b}_{u}" for b, u in QH_SCHEMA + H_SCHEMA]
    variables = {c: {"column": c, "label_de": "", "label_en": "", "unit": ""} for c in canonical_all}
    if var_path.exists():  # keep labels from earlier runs, drop names from older parser versions
        for r in pd.read_csv(var_path, keep_default_na=False).to_dict("records"):
            if r["column"] in variables:
                variables[r["column"]] = r
    new_rows = []

    for y, path in files.items():
        for sheet, freq, sub, schema in ((QH_SHEET, 15, "qh", QH_SCHEMA), (H_SHEET, 60, "vertical_load_1h", H_SCHEMA)):
            target = out_dir / sub / f"{y}.parquet"
            if target.exists() and not a.force:
                print(f"  {y} {sub:17s} exists, skip (use --force)")
                continue
            rec = {"year": y, "file": path.name, "sheet": sheet, "output": str(target.relative_to(root))
                   if target.is_relative_to(root) else str(target), "parsed_at": pd.Timestamp.now().isoformat(timespec="seconds")}
            try:
                rows = read_sheet(path, sheet)
                if rows is None:
                    rec.update(status="sheet_absent")
                    print(f"  {y} {sub:17s} sheet absent")
                    new_rows.append(rec)
                    continue
                df, meta, conv, n_mis, unmapped, missing, unit_mis, src_cols = parse_sheet(rows, freq, schema)
                ts = df["timestamp"]
                y0 = pd.Timestamp(f"{y}-01-01", tz=TZ)
                exp_rows = int((pd.Timestamp(f"{y + 1}-01-01", tz=TZ) - y0) / pd.Timedelta(minutes=freq))
                vals = df[src_cols]
                if ts.iloc[0] != y0:
                    raise ValueError(f"first interval {ts.iloc[0]} is not {y0} (file/year mismatch?)")
                atomic_parquet(df, target)
                for m in meta:
                    variables[m["column"]] = m
                rec.update(
                    status="ok", label_convention=conv, rows=len(df), expected_full_year=exp_rows,
                    complete=len(df) == exp_rows, first_start=str(ts.iloc[0]), last_start=str(ts.iloc[-1]),
                    n_source_columns=len(src_cols), n_missing_columns=len(missing),
                    dst_label_mismatches=n_mis, n_nan=int(vals.isna().sum().sum()),
                    all_nan_columns=";".join(vals.columns[vals.isna().all()]), unmapped_columns=";".join(unmapped),
                    unit_mismatches=" | ".join(unit_mis), missing_columns=";".join(missing),
                )
                flag = "" if rec["complete"] else "  (partial year)"
                print(f"  {y} {sub:17s} ok  {len(df):6d} rows  {len(src_cols)}/{len(schema)} cols  labels={conv}{flag}"
                      + (f"  UNMAPPED: {unmapped}" if unmapped else "")
                      + (f"  UNIT MISMATCH: {len(unit_mis)} col(s), see manifest" if unit_mis else ""))
            except Exception as e:  # one bad file must not stop the run
                rec.update(status="error", error=f"{type(e).__name__}: {e}")
                print(f"  {y} {sub:17s} ERROR {rec['error']}", file=sys.stderr)
            new_rows.append(rec)

    if new_rows:
        new = pd.DataFrame(new_rows)
        if not manifest.empty:
            key = set(zip(new["year"], new["sheet"]))
            manifest = manifest[[(yy, ss) not in key for yy, ss in zip(manifest["year"], manifest["sheet"])]]
        manifest = pd.concat([manifest, new], ignore_index=True).sort_values(["year", "sheet"])
        manifest = manifest.drop(columns=["n_columns"], errors="ignore")  # pre-schema manifest field
        for c in ("rows", "expected_full_year", "n_source_columns", "n_missing_columns", "dst_label_mismatches", "n_nan"):
            if c in manifest:
                manifest[c] = manifest[c].astype("Int64")
        out_dir.mkdir(parents=True, exist_ok=True)
        manifest.to_csv(man_path, index=False)
        pd.DataFrame(list(variables.values())).to_csv(var_path, index=False)
    n_err = sum(r.get("status") == "error" for r in new_rows)
    print(f"Done: {len(new_rows)} sheet(s) processed, {n_err} error(s). Manifest: {man_path}")
    return 1 if n_err else 0


if __name__ == "__main__":
    sys.exit(main())
